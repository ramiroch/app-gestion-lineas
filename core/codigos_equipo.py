"""Códigos de equipo: SIGLA_PLANTA-TIPO-NN (por ejemplo LED-BOE-01, CTX-JEX-A03).

- SIGLA_PLANTA: sigla de la planta (columna "sigla" de la hoja Planta).
- TIPO: abreviatura del tipo de equipo (ABREVIATURAS_TIPO).
- NN: número del equipo con dos dígitos. En plantas con más de una línea va
  adelante la letra de la línea (General → G, Línea A → A, Línea Norte → N).
- Si dos equipos dan el mismo código, el segundo lleva "-2", el tercero "-3", etc.

El código identifica al equipo y nunca cambia.

Migración de una sola vez (agrega la columna "codigo" a data/maestros_clientes.xlsx):
    python -m core.codigos_equipo
"""

import shutil
from collections import Counter, defaultdict
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn

CARPETA_DATOS = Path(__file__).resolve().parent.parent / "data"
RUTA_PLANILLA = CARPETA_DATOS / "maestros_clientes.xlsx"
CARPETA_RESPALDOS = CARPETA_DATOS / "respaldos"

AMARILLO = PatternFill(fill_type="solid", start_color="FFFF00")

ABREVIATURAS_TIPO = {
    "BOE": "BOE",
    "mBOE": "MBOE",
    "JEX": "JEX",
    "Finisher": "FIN",
    "Secador": "SEC",
    "Elevador": "ELE",
    "Alimentador JEX": "ALJ",
    "Calibrador": "CAL",
    "Extractor de prueba": "EXT",
    "Multivane": "MUL",
    "Panel HMI": "HMI",
}


# ---------------------------------------------------------------- armado de códigos

def letra_linea(nombre_linea: str, varias_lineas: bool) -> str:
    """Letra de la línea para el código; vacía si la planta tiene una sola línea.

    "Línea A" → "A", "Línea Norte" → "N", "General" → "G".
    """
    if not varias_lineas:
        return ""
    return nombre_linea.removeprefix("Línea ").strip()[0].upper()


def generar_codigo(sigla: str, tipo: str, numero: int, letra: str = "") -> str:
    """Arma el código SIGLA-TIPO-NN de un equipo. Ejemplo: ("CTX", "JEX", 3, "A") → "CTX-JEX-A03"."""
    if not sigla:
        raise ValueError("Falta la sigla de la planta.")
    if tipo not in ABREVIATURAS_TIPO:
        raise ValueError(f'El tipo "{tipo}" no tiene abreviatura. Agregalo en ABREVIATURAS_TIPO.')
    if not isinstance(numero, int) or numero < 0:
        raise ValueError(f'El número del equipo tiene que ser un entero positivo (dice "{numero}").')
    return f"{sigla}-{ABREVIATURAS_TIPO[tipo]}-{letra}{numero:02d}"


def asignar_codigos(filas: list[dict], siglas: dict) -> list[tuple[str | None, str]]:
    """Arma el código de cada equipo de la lista.

    filas: diccionarios con cliente, planta, linea, tipo y numero.
    siglas: {(cliente, planta): sigla}.
    Devuelve, para cada fila y en el mismo orden, (código o None, nota). La nota
    dice si hubo choque ("choque: ...") o por qué no se pudo armar el código.
    """
    lineas_por_planta = defaultdict(set)
    for fila in filas:
        lineas_por_planta[(fila["cliente"], fila["planta"])].add(fila["linea"])

    usados = Counter()
    resultado = []
    for fila in filas:
        planta = (fila["cliente"], fila["planta"])
        if fila["numero"] is None:
            resultado.append((None, "sin número: no se puede armar el código"))
            continue
        try:
            letra = letra_linea(fila["linea"], len(lineas_por_planta[planta]) > 1)
            codigo = generar_codigo(siglas.get(planta), fila["tipo"], fila["numero"], letra)
        except ValueError as error:
            resultado.append((None, str(error)))
            continue
        usados[codigo] += 1
        if usados[codigo] == 1:
            resultado.append((codigo, ""))
        else:
            resultado.append((f"{codigo}-{usados[codigo]}", f"choque: {codigo} ya estaba usado"))
    return resultado


# ---------------------------------------------------------------- migración de la planilla

def agregar_codigos_a_planilla(ruta: Path = RUTA_PLANILLA,
                               carpeta_respaldos: Path = CARPETA_RESPALDOS) -> list[str]:
    """Agrega la columna "codigo" a la hoja Equipo y pasa PuntoMuestreo a usar el código.

    Hace una copia de respaldo antes de tocar nada. Se niega a correr si la hoja
    Equipo ya tiene códigos (los códigos nunca se regeneran). Devuelve un informe
    con los choques y las filas que quedaron sin código.
    """
    ruta = Path(ruta)
    respaldo = _respaldar(ruta, Path(carpeta_respaldos))
    libro = load_workbook(ruta)
    hoja_equipo, hoja_planta = libro["Equipo"], libro["Planta"]

    columnas_equipo = _encabezado(hoja_equipo, 1)
    if "codigo" in columnas_equipo:
        raise ValueError("La hoja Equipo ya tiene la columna codigo. Los códigos nunca se "
                         "regeneran: si hay que corregir uno, se hace a mano.")

    siglas = _leer_siglas(libro, ruta)
    filas = [dict(zip(columnas_equipo, valores))
             for valores in hoja_equipo.iter_rows(min_row=2, values_only=True)]
    codigos = asignar_codigos(filas, siglas)

    informe = [f"Respaldo: {respaldo}"]
    col_codigo = _agregar_columna(hoja_equipo, 1, "codigo")
    for numero_fila, (fila, (codigo, nota)) in enumerate(zip(filas, codigos), start=2):
        celda = hoja_equipo.cell(numero_fila, col_codigo, codigo)
        if nota:
            celda.fill = AMARILLO
            informe.append(f'Equipo fila {numero_fila} ({fila["planta"]} / {fila["nombre"]}): '
                           f"{codigo or 'SIN CÓDIGO'} – {nota}")

    informe += _puntos_muestreo_a_codigo(libro["PuntoMuestreo"], filas, codigos)
    libro.save(ruta)
    asignados = sum(1 for codigo, _ in codigos if codigo)
    informe.insert(1, f"Códigos asignados: {asignados} de {len(filas)} equipos.")
    return informe


def _respaldar(ruta: Path, carpeta: Path) -> Path:
    """Copia la planilla a la carpeta de respaldos y verifica que la copia exista."""
    carpeta.mkdir(parents=True, exist_ok=True)
    base = f"{ruta.stem}_{date.today():%Y-%m-%d}_antes_de_codigos"
    destino = carpeta / f"{base}{ruta.suffix}"
    numero = 2
    while destino.exists():  # nunca se pisa un respaldo anterior
        destino = carpeta / f"{base}_{numero}{ruta.suffix}"
        numero += 1
    shutil.copy2(ruta, destino)
    if not destino.exists() or destino.stat().st_size != ruta.stat().st_size:
        raise OSError(f"No se pudo hacer el respaldo en {destino}. No se modificó nada.")
    return destino


def _leer_siglas(libro, ruta: Path) -> dict:
    """Lee {(cliente, planta): sigla} de la hoja Planta.

    Si la columna "sigla" no existe, la crea vacía (en amarillo), guarda y da error
    para que se completen las siglas antes de seguir.
    """
    hoja = libro["Planta"]
    columnas = _encabezado(hoja, 1)
    if "sigla" not in columnas:
        col = _agregar_columna(hoja, 1, "sigla")
        for numero_fila in range(2, hoja.max_row + 1):
            hoja.cell(numero_fila, col).fill = AMARILLO
        libro.save(ruta)
        raise ValueError('Agregué la columna "sigla" en la hoja Planta. Completala y volvé a correr.')
    siglas = {}
    for valores in hoja.iter_rows(min_row=2, values_only=True):
        fila = dict(zip(columnas, valores))
        if fila["nombre"] is not None:
            siglas[(fila["cliente"], fila["nombre"])] = fila["sigla"]
    faltan = [planta for planta, sigla in siglas.items() if not sigla]
    if faltan:
        raise ValueError("Faltan siglas en la hoja Planta para: "
                         + ", ".join(nombre for _, nombre in faltan))
    repetidas = [s for s, n in Counter(siglas.values()).items() if n > 1]
    if repetidas:
        raise ValueError(f"Hay siglas repetidas en la hoja Planta: {', '.join(repetidas)}")
    return siglas


def _puntos_muestreo_a_codigo(hoja, filas_equipo: list, codigos: list) -> list[str]:
    """Cambia la columna "equipo (nombre MC)" por "equipo (codigo)" en PuntoMuestreo."""
    fila_enc = next(n for n in range(1, 6) if "equipo (nombre MC)" in _encabezado(hoja, n))
    columnas = _encabezado(hoja, fila_enc)
    col = columnas.index("equipo (nombre MC)") + 1

    codigos_por_nombre = defaultdict(list)
    for fila, (codigo, _) in zip(filas_equipo, codigos):
        codigos_por_nombre[(fila["cliente"], fila["planta"], fila["nombre"])].append(codigo)

    informe = []
    for numero_fila in range(fila_enc + 1, hoja.max_row + 1):
        valores = dict(zip(columnas, (c.value for c in hoja[numero_fila])))
        if valores.get("equipo (nombre MC)") is None:
            continue
        candidatos = codigos_por_nombre[(valores["cliente"], valores["planta"],
                                         valores["equipo (nombre MC)"])]
        celda = hoja.cell(numero_fila, col)
        if len(candidatos) == 1 and candidatos[0]:
            celda.value = candidatos[0]
        else:
            celda.fill = AMARILLO
            informe.append(f"PuntoMuestreo fila {numero_fila}: no se pudo pasar "
                           f'"{celda.value}" a código ({len(candidatos)} equipos con ese nombre).')
    _renombrar_columna(hoja, fila_enc, col, "equipo (codigo)")
    return informe


# ---------------------------------------------------------------- ayudas para tablas de Excel

def _encabezado(hoja, numero_fila: int) -> list:
    return [celda.value for celda in hoja[numero_fila]]


def _agregar_columna(hoja, fila_encabezado: int, nombre: str) -> int:
    """Agrega una columna al final, con el estilo del encabezado vecino, y la suma a la tabla."""
    col = hoja.max_column + 1
    celda = hoja.cell(fila_encabezado, col, nombre)
    celda._style = copy(hoja.cell(fila_encabezado, col - 1)._style)
    hoja.column_dimensions[get_column_letter(col)].width = 16
    for tabla in hoja.tables.values():
        col_ini, fila_ini, col_fin, fila_fin = range_boundaries(tabla.ref)
        if col_fin == col - 1:
            tabla.ref = f"{get_column_letter(col_ini)}{fila_ini}:{get_column_letter(col)}{fila_fin}"
            if tabla.autoFilter is not None:
                tabla.autoFilter.ref = tabla.ref
            tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=nombre))
    return col


def _renombrar_columna(hoja, fila_encabezado: int, col: int, nombre: str) -> None:
    """Cambia el nombre de una columna en la celda y en la definición de la tabla."""
    anterior = hoja.cell(fila_encabezado, col).value
    hoja.cell(fila_encabezado, col).value = nombre
    for tabla in hoja.tables.values():
        for columna in tabla.tableColumns:
            if columna.name == anterior:
                columna.name = nombre


if __name__ == "__main__":
    for linea in agregar_codigos_a_planilla():
        print(linea)
