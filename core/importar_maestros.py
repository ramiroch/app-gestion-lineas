"""Importa los datos maestros desde data/maestros_clientes.xlsx a la base.

Lee las hojas Cliente, Planta, Linea, Equipo, PuntoMuestreo y RangoReferencia
(en ese orden, porque cada una necesita a su "padre" ya cargado).

- Se puede correr muchas veces: si un registro ya existe se actualiza, no se duplica.
- Una celda vacía se guarda vacía (también los parámetros de cálculo de la planta).
- Las filas con error no se cargan y se anotan en data/reporte_maestros.csv.
- Un equipo se reconoce solo por su código (columna "codigo", que nunca cambia).
  PuntoMuestreo busca el equipo por ese código. Si en la base queda un equipo cuyo
  código ya no está en la planilla, se avisa en el reporte.

Uso (desde la carpeta del proyecto):  python -m core.importar_maestros
"""

import csv
from collections import Counter
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from core.crear_base import crear_tablas
from core.db import crear_motor, crear_sesion
from core.models import (ESTADOS_EQUIPO, Cliente, Equipo, Linea, Planta, PuntoMuestreo,
                         RangoReferencia)

CARPETA_DATOS = Path(__file__).resolve().parent.parent / "data"
RUTA_PLANILLA = CARPETA_DATOS / "maestros_clientes.xlsx"
RUTA_REPORTE = CARPETA_DATOS / "reporte_maestros.csv"

# Columnas que tiene que tener cada hoja. Las demás (por ejemplo "observaciones") se ignoran.
COLUMNAS = {
    "Cliente": ["nombre", "pais_casa_matriz"],
    "Planta": ["cliente", "nombre", "pais", "blanco_ml", "factor_aceite_fruta",
               "factor_aceite_emulsion", "hora_inicio_dia"],
    "Linea": ["cliente", "planta", "nombre"],
    "Equipo": ["codigo", "cliente", "planta", "linea", "nombre", "tipo", "numero", "modelo",
               "serie_mc", "numero_serie", "codigo_activo_mc", "numero_activo_fijo", "estado"],
    "PuntoMuestreo": ["cliente", "planta", "equipo (codigo)", "tipo", "nombre"],
    "RangoReferencia": ["planta", "variable", "punto o equipo", "fruta", "minimo", "maximo",
                        "unidad", "fuente"],
}

# Columnas que identifican un registro: dos filas con los mismos valores son el mismo registro.
CLAVES = {
    "Cliente": ["nombre"],
    "Planta": ["cliente", "nombre"],
    "Linea": ["cliente", "planta", "nombre"],
    "Equipo": ["codigo"],
    "PuntoMuestreo": ["cliente", "planta", "equipo (codigo)", "nombre"],
    "RangoReferencia": ["planta", "variable", "punto o equipo", "fruta"],
}


class ErrorFila(Exception):
    """Problema en una fila de la planilla. El mensaje es el motivo que va al reporte."""


# ---------------------------------------------------------------- lectura de la planilla

def leer_hoja(libro, hoja: str) -> list[tuple[int, dict]]:
    """Devuelve las filas con datos de una hoja como (n.º de fila de Excel, {columna: valor}).

    El encabezado se busca en las primeras 5 filas, porque algunas hojas tienen
    una nota arriba. Los textos se recortan y las celdas vacías quedan como None.
    """
    if hoja not in libro.sheetnames:
        raise ValueError(f'La planilla no tiene la hoja "{hoja}".')
    ws = libro[hoja]
    fila_encabezado = _buscar_encabezado(ws, hoja)
    encabezado = [_limpiar(celda.value) for celda in ws[fila_encabezado]]
    filas = []
    for numero_fila, valores in enumerate(
            ws.iter_rows(min_row=fila_encabezado + 1, values_only=True), start=fila_encabezado + 1):
        datos = {columna: _limpiar(valor) for columna, valor in zip(encabezado, valores) if columna}
        if any(valor is not None for valor in datos.values()):
            filas.append((numero_fila, datos))
    return filas


def _buscar_encabezado(ws, hoja: str) -> int:
    """Devuelve el número de la fila que tiene todas las columnas esperadas de la hoja."""
    for numero_fila in range(1, 6):
        encabezado = [_limpiar(celda.value) for celda in ws[numero_fila]]
        if all(columna in encabezado for columna in COLUMNAS[hoja]):
            return numero_fila
    raise ValueError(f'La hoja "{hoja}" no tiene las columnas: {", ".join(COLUMNAS[hoja])}.')


def _limpiar(valor):
    """Recorta los espacios de los textos y convierte las celdas vacías en None."""
    if isinstance(valor, str):
        valor = valor.strip()
        return valor or None
    return valor


# ---------------------------------------------------------------- conversión de valores

def _obligatorios(datos: dict, *columnas: str) -> None:
    """Da error si alguna de las columnas está vacía (las nombra todas)."""
    faltan = [columna for columna in columnas if datos.get(columna) is None]
    if faltan:
        raise ErrorFila("Falta " + ", ".join(f'"{columna}"' for columna in faltan) + ".")


def _texto(valor) -> str | None:
    """Convierte a texto (un n.º de serie puede venir como número desde Excel).

    Los números enteros guardados como decimales (1.0) se escriben sin decimales ("1").
    """
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return None if valor is None else str(valor)


def _numero(datos: dict, columna: str) -> float | None:
    """Convierte a número decimal; acepta coma o punto como separador."""
    valor = datos.get(columna)
    if valor is None or isinstance(valor, (int, float)):
        return valor
    try:
        return float(str(valor).replace(",", "."))
    except ValueError:
        raise ErrorFila(f'"{columna}" tiene que ser un número (dice "{valor}").')


def _entero(datos: dict, columna: str) -> int | None:
    """Convierte a número entero (2 o 2.0 sirven; 2.5 o "dos" no)."""
    valor = _numero(datos, columna)
    if valor is None:
        return None
    if float(valor).is_integer():
        return int(valor)
    raise ErrorFila(f'"{columna}" tiene que ser un número entero (dice "{valor}").')


def _hora(datos: dict, columna: str) -> time | None:
    """Convierte a hora. Acepta una hora de Excel o un texto como "05:00"."""
    valor = datos.get(columna)
    if valor is None or isinstance(valor, time):
        return valor
    if isinstance(valor, datetime):
        return valor.time()
    try:
        return time.fromisoformat(str(valor))
    except ValueError:
        raise ErrorFila(f'"{columna}" tiene que ser una hora como 05:00 (dice "{valor}").')


def _defecto(modelo, columna: str):
    """Valor por defecto de una columna, tal como está definido en core/models.py."""
    return modelo.__table__.columns[columna].default.arg


# ---------------------------------------------------------------- búsquedas

def _buscar(lista, **condiciones):
    """Devuelve el primer objeto de la lista cuyos campos coinciden con las condiciones."""
    for objeto in lista:
        if all(getattr(objeto, campo) == valor for campo, valor in condiciones.items()):
            return objeto
    return None


def _buscar_cliente(sesion: Session, nombre: str) -> Cliente:
    cliente = sesion.scalars(select(Cliente).where(Cliente.nombre == nombre)).first()
    if cliente is None:
        raise ErrorFila(f'El cliente "{nombre}" no existe en la hoja Cliente (o su fila tuvo un error).')
    return cliente


def _buscar_planta(sesion: Session, nombre_cliente: str, nombre_planta: str) -> Planta:
    cliente = _buscar_cliente(sesion, nombre_cliente)
    planta = _buscar(cliente.plantas, nombre=nombre_planta)
    if planta is None:
        raise ErrorFila(f'La planta "{nombre_planta}" del cliente "{nombre_cliente}" no existe '
                        "en la hoja Planta (o su fila tuvo un error).")
    return planta


# ---------------------------------------------------------------- una función por hoja
# Cada una valida la fila y devuelve (registro existente o None, campos a guardar).

def _fila_cliente(sesion: Session, datos: dict):
    _obligatorios(datos, "nombre", "pais_casa_matriz")
    campos = {"nombre": _texto(datos["nombre"]), "pais": _texto(datos["pais_casa_matriz"])}
    existente = sesion.scalars(select(Cliente).where(Cliente.nombre == campos["nombre"])).first()
    return existente, campos


def _fila_planta(sesion: Session, datos: dict):
    _obligatorios(datos, "cliente", "nombre", "pais")
    cliente = _buscar_cliente(sesion, datos["cliente"])
    campos = {
        "cliente": cliente,
        "nombre": _texto(datos["nombre"]),
        "pais": _texto(datos["pais"]),
        "blanco_ml": _numero(datos, "blanco_ml"),
        "factor_aceite_fruta": _numero(datos, "factor_aceite_fruta"),
        "factor_aceite_emulsion": _numero(datos, "factor_aceite_emulsion"),
        "hora_inicio_dia": _hora(datos, "hora_inicio_dia"),
    }
    return _buscar(cliente.plantas, nombre=campos["nombre"]), campos


def _fila_linea(sesion: Session, datos: dict):
    _obligatorios(datos, "cliente", "planta", "nombre")
    planta = _buscar_planta(sesion, datos["cliente"], datos["planta"])
    campos = {"planta": planta, "nombre": _texto(datos["nombre"])}
    return _buscar(planta.lineas, nombre=campos["nombre"]), campos


def _fila_equipo(sesion: Session, datos: dict):
    _obligatorios(datos, "codigo", "cliente", "planta", "linea", "nombre", "tipo", "numero")
    planta = _buscar_planta(sesion, datos["cliente"], datos["planta"])
    linea = _buscar(planta.lineas, nombre=datos["linea"])
    if linea is None:
        raise ErrorFila(f'La línea "{datos["linea"]}" de la planta "{datos["planta"]}" no existe '
                        "en la hoja Linea (o su fila tuvo un error).")
    estado = datos["estado"] or _defecto(Equipo, "estado")
    if estado not in ESTADOS_EQUIPO:
        raise ErrorFila(f'Estado "{estado}" no válido. Usá uno de: {", ".join(ESTADOS_EQUIPO)}.')
    campos = {
        "codigo": _texto(datos["codigo"]),
        "linea": linea,
        "nombre": _texto(datos["nombre"]),
        "tipo": _texto(datos["tipo"]),
        "numero": _entero(datos, "numero"),
        "estado": estado,
        "modelo": _texto(datos["modelo"]),
        "serie_mc": _texto(datos["serie_mc"]),
        "numero_serie": _texto(datos["numero_serie"]),
        "codigo_activo_mc": _texto(datos["codigo_activo_mc"]),
        "numero_activo_fijo": _texto(datos["numero_activo_fijo"]),
    }
    existente = sesion.scalars(select(Equipo).where(Equipo.codigo == campos["codigo"])).first()
    return existente, campos


def _fila_punto_muestreo(sesion: Session, datos: dict):
    _obligatorios(datos, "cliente", "planta", "equipo (codigo)", "tipo", "nombre")
    planta = _buscar_planta(sesion, datos["cliente"], datos["planta"])
    codigo = _texto(datos["equipo (codigo)"])
    equipo = sesion.scalars(select(Equipo).where(Equipo.codigo == codigo)).first()
    if equipo is None:
        raise ErrorFila(f'El equipo "{codigo}" no existe en la hoja Equipo (o su fila tuvo un error).')
    if equipo.linea.planta is not planta:
        raise ErrorFila(f'El equipo "{codigo}" es de la planta "{equipo.linea.planta.nombre}", '
                        f'no de "{datos["planta"]}".')
    campos = {"equipo": equipo, "tipo": _texto(datos["tipo"]), "nombre": _texto(datos["nombre"])}
    return _buscar(equipo.puntos_muestreo, nombre=campos["nombre"]), campos


def _fila_rango(sesion: Session, datos: dict):
    _obligatorios(datos, "planta", "variable")
    plantas = sesion.scalars(select(Planta).where(Planta.nombre == datos["planta"])).all()
    if not plantas:
        raise ErrorFila(f'La planta "{datos["planta"]}" no existe en la hoja Planta '
                        "(o su fila tuvo un error).")
    if len(plantas) > 1:
        raise ErrorFila(f'Hay más de una planta llamada "{datos["planta"]}": no se sabe cuál es.')
    campos = {
        "planta": plantas[0],
        "variable": _texto(datos["variable"]),
        "punto_o_equipo": _texto(datos["punto o equipo"]),
        "fruta": _texto(datos["fruta"]),
        "minimo": _numero(datos, "minimo"),
        "maximo": _numero(datos, "maximo"),
        "unidad": _texto(datos["unidad"]),
        "fuente": _texto(datos["fuente"]),
    }
    if campos["minimo"] is not None and campos["maximo"] is not None \
            and campos["minimo"] > campos["maximo"]:
        raise ErrorFila(f'El mínimo ({campos["minimo"]}) es mayor que el máximo ({campos["maximo"]}).')
    existente = _buscar(plantas[0].rangos, variable=campos["variable"],
                        punto_o_equipo=campos["punto_o_equipo"], fruta=campos["fruta"])
    return existente, campos


# Orden de importación: cada hoja necesita que su "padre" ya esté cargado.
HOJAS = [
    ("Cliente", Cliente, _fila_cliente),
    ("Planta", Planta, _fila_planta),
    ("Linea", Linea, _fila_linea),
    ("Equipo", Equipo, _fila_equipo),
    ("PuntoMuestreo", PuntoMuestreo, _fila_punto_muestreo),
    ("RangoReferencia", RangoReferencia, _fila_rango),
]


# ---------------------------------------------------------------- importación

def importar_maestros(sesion: Session, ruta: Path = RUTA_PLANILLA) -> tuple[dict, list, list]:
    """Carga la planilla de datos maestros en la base.

    Devuelve (resumen, errores, avisos):
    - resumen: {hoja: {"creados": n, "actualizados": n}}
    - errores: lista de (hoja, n.º de fila de Excel, motivo) de las filas que no se cargaron.
    - avisos: lista de (hoja, "", motivo) de equipos que están en la base pero no en la planilla.
    """
    libro = load_workbook(ruta, data_only=True)
    resumen = {}
    errores = []
    avisos = []
    for hoja, modelo, procesar_fila in HOJAS:
        resumen[hoja] = {"creados": 0, "actualizados": 0}
        filas = leer_hoja(libro, hoja)
        repeticiones = Counter(_clave(hoja, datos) for _, datos in filas)
        for numero_fila, datos in filas:
            motivos = []
            clave = _clave(hoja, datos)
            veces = repeticiones[clave]
            if veces > 1 and any(valor is not None for valor in clave):  # vacía: ya la reporta "Falta"
                motivos.append(_motivo_repetido(hoja, datos, veces))
            try:
                existente, campos = procesar_fila(sesion, datos)
            except ErrorFila as error:
                motivos.append(str(error))
            if motivos:
                errores.append((hoja, numero_fila, " · ".join(motivos)))
                continue
            if existente is None:
                sesion.add(modelo(**campos))
                resumen[hoja]["creados"] += 1
            else:
                for campo, valor in campos.items():
                    setattr(existente, campo, valor)
                resumen[hoja]["actualizados"] += 1
            sesion.flush()
        if hoja == "Equipo":
            avisos = _equipos_sobrantes(sesion, filas)
    sesion.commit()
    return resumen, errores, avisos


def _clave(hoja: str, datos: dict) -> tuple:
    return tuple(_texto(datos.get(columna)) for columna in CLAVES[hoja])


def _equipos_sobrantes(sesion: Session, filas: list) -> list:
    """Avisa de los equipos que están en la base pero no en la hoja Equipo (no los borra)."""
    codigos_planilla = {_texto(datos.get("codigo")) for _, datos in filas}
    avisos = []
    for equipo in sesion.scalars(select(Equipo)).all():
        if equipo.codigo not in codigos_planilla:
            avisos.append(("Equipo", "", f'Aviso: el equipo {equipo.codigo} ("{equipo.nombre}", '
                           f'planta "{equipo.linea.planta.nombre}") está en la base pero no en la planilla.'))
    return avisos


def _motivo_repetido(hoja: str, datos: dict, veces: int) -> str:
    descripcion = ", ".join(
        f"{columna} vacío" if datos.get(columna) is None else f'{columna} "{datos.get(columna)}"'
        for columna in CLAVES[hoja])
    distinguir = " o ".join(c for c in CLAVES[hoja] if c not in ("cliente", "planta"))
    return f"Hay {veces} filas con el mismo {descripcion}. Diferencialas por {distinguir}."


# ---------------------------------------------------------------- reporte y resumen

def guardar_reporte(errores: list, ruta: Path = RUTA_REPORTE) -> Path:
    """Escribe los errores en un CSV (hoja; fila; motivo) que se abre directo en Excel.

    Si no hay errores, el archivo queda solo con el encabezado.
    """
    ruta = Path(ruta)
    ruta.parent.mkdir(exist_ok=True)
    with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.writer(archivo, delimiter=";")
        escritor.writerow(["hoja", "fila", "motivo"])
        escritor.writerows(errores)
    return ruta


def texto_resumen(resumen: dict, errores: list, avisos: list,
                  ruta_reporte: Path = RUTA_REPORTE) -> str:
    """Arma el resumen para mostrar en pantalla."""
    lineas = ["Importación de datos maestros"]
    for hoja, cuentas in resumen.items():
        lineas.append(f"  {hoja + ':':<17}{cuentas['creados']:>4} creados, "
                      f"{cuentas['actualizados']:>4} actualizados")
    lineas.append(f"  Errores: {len(errores)} · Avisos: {len(avisos)}"
                  + (f" (detalle en {ruta_reporte})" if errores or avisos else ""))
    return "\n".join(lineas)


if __name__ == "__main__":
    motor = crear_motor()
    crear_tablas(motor)  # crea las tablas que falten (por ejemplo rango_referencia)
    with crear_sesion(motor) as sesion:
        resumen, errores, avisos = importar_maestros(sesion)
    guardar_reporte(errores + avisos)
    print(texto_resumen(resumen, errores, avisos))
