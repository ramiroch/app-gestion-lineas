"""Pruebas de core/codigos_equipo.py con una planilla chica inventada (no la real)."""

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from core.codigos_equipo import (agregar_codigos_a_planilla, asignar_codigos, generar_codigo,
                                 letra_linea)


def test_generar_codigo():
    assert generar_codigo("LED", "BOE", 1) == "LED-BOE-01"
    assert generar_codigo("CTX", "JEX", 3, "A") == "CTX-JEX-A03"
    assert generar_codigo("CVA", "Alimentador JEX", 13) == "CVA-ALJ-13"


def test_generar_codigo_valida_entradas():
    with pytest.raises(ValueError, match="no tiene abreviatura"):
        generar_codigo("LED", "Tostadora", 1)
    with pytest.raises(ValueError, match="sigla"):
        generar_codigo(None, "BOE", 1)
    with pytest.raises(ValueError, match="entero"):
        generar_codigo("LED", "BOE", 1.5)


def test_letra_linea():
    assert letra_linea("Línea 1", varias_lineas=False) == ""
    assert letra_linea("Línea A", varias_lineas=True) == "A"
    assert letra_linea("Línea Norte", varias_lineas=True) == "N"
    assert letra_linea("General", varias_lineas=True) == "G"


def equipo(planta, linea, tipo, numero):
    return {"cliente": "CLI", "planta": planta, "linea": linea, "tipo": tipo, "numero": numero}


def test_asignar_codigos_con_choques_y_sin_numero():
    filas = [
        equipo("Norte", "Línea 1", "JEX", 1),
        equipo("Norte", "Línea 1", "JEX", 1),     # choque
        equipo("Norte", "Línea 1", "JEX", 1),     # segundo choque
        equipo("Norte", "Línea 1", "JEX", None),  # sin número
        equipo("Sur", "Línea A", "BOE", 2),
        equipo("Sur", "Línea B", "BOE", 2),       # otra línea: no choca
    ]
    siglas = {("CLI", "Norte"): "NOR", ("CLI", "Sur"): "SUR"}

    codigos = [codigo for codigo, _ in asignar_codigos(filas, siglas)]

    assert codigos == ["NOR-JEX-01", "NOR-JEX-01-2", "NOR-JEX-01-3", None, "SUR-BOE-A02", "SUR-BOE-B02"]


def crear_planilla(ruta, con_sigla=True):
    """Planilla mínima con las hojas Planta, Equipo y PuntoMuestreo, cada una con su tabla."""
    libro = Workbook()
    libro.remove(libro.active)
    planta = libro.create_sheet("Planta")
    planta.append(["cliente", "nombre"] + (["sigla"] if con_sigla else []))
    planta.append(["CLI", "Norte"] + (["NOR"] if con_sigla else []))
    planta.add_table(Table(displayName="T_Planta", ref=f"A1:{'C' if con_sigla else 'B'}2"))
    equipos = libro.create_sheet("Equipo")
    equipos.append(["cliente", "planta", "linea", "nombre", "tipo", "numero"])
    equipos.append(["CLI", "Norte", "Línea 1", "BOE N°1", "BOE", 1])
    equipos.append(["CLI", "Norte", "Línea 1", "JEX", "JEX", 1])
    equipos.append(["CLI", "Norte", "Línea 1", "JEX", "JEX", 1])
    equipos.add_table(Table(displayName="T_Equipo", ref="A1:F4"))
    puntos = libro.create_sheet("PuntoMuestreo")
    puntos.append(["Nota"])
    puntos.append([])
    puntos.append(["cliente", "planta", "equipo (nombre MC)", "tipo", "nombre"])
    puntos.append(["CLI", "Norte", "BOE N°1", "fruta entrada", "BOE 1 – entrada"])
    puntos.append(["CLI", "Norte", "JEX", "emulsión", "JEX – emulsión"])
    puntos.add_table(Table(displayName="T_PuntoMuestreo", ref="A3:E5"))
    libro.save(ruta)
    return ruta


def test_migracion_agrega_codigos_y_respaldo(tmp_path):
    ruta = crear_planilla(tmp_path / "maestros.xlsx")

    informe = agregar_codigos_a_planilla(ruta, tmp_path / "respaldos")

    assert len(list((tmp_path / "respaldos").iterdir())) == 1
    libro = load_workbook(ruta)
    equipos = libro["Equipo"]
    assert [c.value for c in equipos["G"]] == ["codigo", "NOR-BOE-01", "NOR-JEX-01", "NOR-JEX-01-2"]
    assert equipos.tables["T_Equipo"].ref == "A1:G4"
    assert equipos["G4"].fill.start_color.rgb.endswith("FFFF00")   # el choque queda en amarillo
    assert any("choque" in linea for linea in informe)

    puntos = libro["PuntoMuestreo"]
    assert puntos["C3"].value == "equipo (codigo)"
    assert puntos["C4"].value == "NOR-BOE-01"
    assert puntos["C5"].value == "JEX"                              # nombre ambiguo: queda igual
    assert any("PuntoMuestreo fila 5" in linea for linea in informe)
    assert [c.name for c in puntos.tables["T_PuntoMuestreo"].tableColumns][2] == "equipo (codigo)"


def test_migracion_no_corre_dos_veces(tmp_path):
    ruta = crear_planilla(tmp_path / "maestros.xlsx")
    agregar_codigos_a_planilla(ruta, tmp_path / "respaldos")

    with pytest.raises(ValueError, match="ya tiene la columna codigo"):
        agregar_codigos_a_planilla(ruta, tmp_path / "respaldos")
    assert len(list((tmp_path / "respaldos").iterdir())) == 2      # no pisa el respaldo anterior


def test_migracion_sin_siglas_agrega_la_columna_y_se_detiene(tmp_path):
    ruta = crear_planilla(tmp_path / "maestros.xlsx", con_sigla=False)

    with pytest.raises(ValueError, match="Completala"):
        agregar_codigos_a_planilla(ruta, tmp_path / "respaldos")

    libro = load_workbook(ruta)
    assert libro["Planta"]["C1"].value == "sigla"
    assert "codigo" not in [c.value for c in libro["Equipo"][1]]
